#!/usr/bin/env python3
"""
parse_check.py — Bir CV dosyasinin ATS ayristiricisi acisindan okunabilirligini test eder.

Ne yapar: dosyadan metni geri cikarir (parser'in gordugu sey budur) ve su kontrolleri raporlar:
  1. Secilebilir metin var mi / gorsel PDF mi
  2. Iletisim bilgisi bulunabiliyor mu (e-posta, telefon, LinkedIn)
  3. Standart bolum basliklari taninabiliyor mu (TR + EN)
  4. Tarih formati tutarli mi
  5. Sutun / tablo / metin kutusu / header-footer suphesi
  6. Kanit yogunlugu: rakam veya kapsam iceren madde orani
  7. Klise ve yapay-zeka kalibi ifade taramasi
  8. (--job-description verilirse) ilandaki tekrar eden terimlerin karsilanma orani

Kullanim:
    python3 parse_check.py cv.pdf
    python3 parse_check.py cv.docx --job-description ilan.txt
    python3 parse_check.py cv.pdf --json
"""

import argparse
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

SECTIONS = {
    "deneyim": ["deneyim", "iş deneyimi", "is deneyimi", "tecrübe", "experience",
                "work experience", "professional experience", "employment"],
    "eğitim": ["eğitim", "egitim", "öğrenim", "education", "academic"],
    "beceriler": ["beceri", "yetkinlik", "skills", "technical skills", "core competencies"],
    "özet": ["özet", "ozet", "profil", "hakkımda", "summary", "profile", "about"],
}

CLICHES = [
    "sonuç odaklı", "sonuc odakli", "takım oyuncusu", "takim oyuncusu", "detay odaklı",
    "detay odakli", "hızlı öğrenen", "hizli ogrenen", "güçlü iletişim becerileri",
    "proaktif", "dinamik bir", "katma değer", "katma deger",
    "results-driven", "results driven", "passionate about", "team player",
    "detail-oriented", "self-starter", "go-getter", "think outside the box",
    "synergy", "best-in-class", "hard worker", "excellent communication skills",
]

AI_TELLS = ["spearheaded", "leveraged", "orchestrated", "championed", "revolutionized",
            "seamlessly", "robust suite", "cutting-edge", "game-changing"]

STOPWORDS = set("""
ve veya ile için icin bir bu şu su da de den dan olarak gibi daha en çok cok ama ancak
the a an and or for with to of in on at as by from is are be will you your we our their
that this those these it its not have has had can may must should would could our you
job role work team company position candidate experience experiences years year new
""".split())


# ---------- metin cikarma ----------

def extract_pdf(path):
    pages = []
    try:
        out = subprocess.run(["pdftotext", "-layout", str(path), "-"],
                             capture_output=True, text=True, timeout=60)
        if out.returncode == 0:
            pages = out.stdout.split("\f")
            while pages and not pages[-1].strip():
                pages.pop()   # pdftotext sonda bos parca birakir
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    if not any(p.strip() for p in pages):
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            pages = [(p.extract_text() or "") for p in reader.pages]
        except Exception as e:
            return [], {"hata": f"PDF okunamadi: {e}"}
    return pages, {}


def extract_docx(path):
    meta = {"tablo_sayisi": 0, "metin_kutusu": False, "header_footer_metni": ""}
    try:
        from docx import Document
        doc = Document(str(path))
        paras = [p.text for p in doc.paragraphs]
        meta["tablo_sayisi"] = len(doc.tables)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    paras.append(cell.text)
        hf = []
        for sec in doc.sections:
            for part in (sec.header, sec.footer):
                hf += [p.text for p in part.paragraphs if p.text.strip()]
        meta["header_footer_metni"] = " | ".join(hf)
    except Exception as e:
        return [], {"hata": f"DOCX okunamadi: {e}"}
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
            meta["metin_kutusu"] = ("txbxContent" in xml) or ("v:textbox" in xml)
    except Exception:
        pass
    return ["\n".join(paras)], meta


def extract(path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in (".txt", ".md"):
        return [path.read_text(encoding="utf-8", errors="ignore")], {}
    return [], {"hata": f"Desteklenmeyen dosya turu: {ext} (pdf, docx, txt, md)"}


# ---------- kontroller ----------

def check_text_presence(pages, meta):
    findings = []
    total = sum(len(p.strip()) for p in pages)
    if meta.get("hata"):
        findings.append(("KRITIK", meta["hata"]))
        return findings, total
    if total < 200:
        findings.append(("KRITIK",
            "Dosyadan neredeyse hic metin cikmadi. Buyuk ihtimalle gorsel/taranmis PDF — "
            "ayristirici hicbir sey okuyamaz. Metin tabanli olarak yeniden uret."))
    else:
        bos = [i + 1 for i, p in enumerate(pages) if len(p.strip()) < 50]
        if bos and len(pages) > 1:
            findings.append(("ONEMLI", f"Su sayfalardan metin cikmadi: {bos} — gorsel olabilir."))
    return findings, total


def check_contact(text):
    f = []
    email = re.search(r"[\w\.\-\+]+@[\w\-]+\.[\w\.\-]+", text)
    phone = re.search(r"(\+?\d[\d\s\-\(\)]{8,}\d)", text)
    linkedin = re.search(r"linkedin\.com/[\w\-/]+", text, re.I)
    if not email:
        f.append(("KRITIK", "E-posta adresi metinden cikarilamadi. Gorsel/ikon icine gomulmus olabilir."))
    if not phone:
        f.append(("ONEMLI", "Telefon numarasi metinden cikarilamadi."))
    if not linkedin:
        f.append(("INCE", "LinkedIn baglantisi bulunamadi — recruiter'in ilk baktigi yerdir."))
    return f, {"email": bool(email), "telefon": bool(phone), "linkedin": bool(linkedin)}


def check_sections(text):
    low = text.lower()
    found, missing = [], []
    for key, variants in SECTIONS.items():
        if any(v in low for v in variants):
            found.append(key)
        else:
            missing.append(key)
    f = []
    for m in missing:
        sev = "ONEMLI" if m in ("deneyim", "eğitim") else "INCE"
        f.append((sev, f"'{m}' bolum basligi taninamadi. Standart baslik kullan "
                       f"(orn. {SECTIONS[m][0].title()} / {SECTIONS[m][-1].title()})."))
    return f, found


MONTH_YEAR = (r"\b(?:Oca|Şub|Sub|Mar|Nis|May|Haz|Tem|Ağu|Agu|Eyl|Eki|Kas|Ara|"
              r"Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-zçğıöşü]*\.?\s+\d{4}\b")


def check_dates(text):
    numeric = re.findall(r"\b\d{1,2}[/.]\d{4}\b", text)
    monthly = re.findall(MONTH_YEAR, text, re.I)
    # "Tem 2024" icindeki yil, ciplak yil olarak ikinci kez sayilmasin
    residue = re.sub(MONTH_YEAR, " ", text, flags=re.I)
    residue = re.sub(r"\b\d{1,2}[/.]\d{4}\b", " ", residue)
    bare = re.findall(r"(?<![\d/.])\b(?:19|20)\d{2}\b(?![\d/.])", residue)
    patterns = {"MM/YYYY": len(numeric), "Ay YYYY": len(monthly), "YYYY": len(bare)}
    f = []
    rol_formatlari = {k: v for k, v in patterns.items() if k != "YYYY" and v >= 2}
    if len(rol_formatlari) > 1:
        f.append(("ONEMLI", "Rol tarihlerinde iki farkli format karisik: "
                            + ", ".join(f"{k} ({v})" for k, v in rol_formatlari.items())
                            + ". Tek formata indir."))
    elif rol_formatlari and patterns["YYYY"] >= 2:
        f.append(("INCE", f"Ciplak yil kullanimi da var ({patterns['YYYY']} adet). "
                          "Egitim ve sertifikada yalnizca yil normaldir; rol tarihlerinde degil."))
    if not any(v >= 2 for v in patterns.values()):
        f.append(("KRITIK", "Tarih araligi bulunamadi. Ayristirici rol surelerini cikaramaz."))
    return f, patterns


def check_layout(pages, meta, ext):
    f = []
    if ext == ".docx":
        if meta.get("tablo_sayisi"):
            f.append(("ONEMLI", f"{meta['tablo_sayisi']} tablo var. Tablo icindeki metin "
                                "ayristirmada sik sik karisir; tek sutun akisa cevir."))
        if meta.get("metin_kutusu"):
            f.append(("KRITIK", "Belgede metin kutusu var. Cogu ayristirici metin kutusunu atlar."))
        if meta.get("header_footer_metni"):
            snippet = meta["header_footer_metni"][:80]
            if re.search(r"@|\d{3}", snippet):
                f.append(("KRITIK", f"Header/footer'da iletisim gibi bilgi var ('{snippet}'). "
                                    "Ayristirici burayi cogu zaman okumaz — govdeye tasi."))
            else:
                f.append(("INCE", "Header/footer'da metin var; kritik bilgi koyma."))
    else:
        # PDF: satir ortasinda genis bosluk = sutun/tablo suphesi
        suspicious = 0
        lines = 0
        for p in pages:
            for line in p.splitlines():
                s = line.rstrip()
                if len(s) < 25:
                    continue
                lines += 1
                if re.search(r"\S {4,}\S", s.strip()):
                    suspicious += 1
        if lines and suspicious / lines > 0.25:
            f.append(("ONEMLI", f"Satirlarin %{suspicious/lines*100:.0f}'inde genis ic bosluk var — "
                                "iki sutunlu yerlesim veya tablo olabilir. Okuma sirasi bozulabilir; "
                                "cikan metni gozle kontrol et."))
    return f


def check_evidence(text):
    lines = [l.strip(" •-–\t") for l in text.splitlines()]
    bullets = [l for l in lines if 40 <= len(l) <= 400]
    if not bullets:
        return [("ONEMLI", "Madde benzeri satir bulunamadi; icerik cok kisa veya biçim bozuk.")], {}
    quant = [b for b in bullets if re.search(r"\d", b)]
    oran = len(quant) / len(bullets)
    f = []
    if oran < 0.30:
        f.append(("KRITIK", f"Rakam iceren madde orani cok dusuk: %{oran*100:.0f}. "
                            "Sorumluluk anlatiyorsun, sonuc degil."))
    elif oran < 0.45:
        f.append(("ONEMLI", f"Rakam iceren madde orani %{oran*100:.0f}. Hedef %45+."))
    uzun = [b for b in bullets if len(b) > 220]
    if uzun:
        f.append(("INCE", f"{len(uzun)} paragraf/madde iki satiri asiyor (ozet paragrafi da sayilmis olabilir); ilk taramada okunmaz."))
    return f, {"madde": len(bullets), "rakamli": len(quant), "oran": round(oran, 2)}


def check_language(text):
    low = text.lower()
    f = []
    hits = sorted({c for c in CLICHES if c in low})
    if hits:
        f.append(("ONEMLI", f"Icı bos klise ifadeler: {', '.join(hits[:6])}"
                            f"{' ...' if len(hits) > 6 else ''}. Yerine kanit yaz."))
    ai = sorted({a for a in AI_TELLS if a in low})
    if len(ai) >= 3:
        f.append(("ONEMLI", f"Asiri kullanilan fiil/sifat yigilmasi: {', '.join(ai)}. "
                            "Insan recruiter bunu yapay zeka kalibi olarak okuyor; cesitlendir."))
    if re.search(r"\b(i̇|i)ş\s?bu\b|saygilarimla|sincerely", low):
        f.append(("INCE", "CV icinde mektup dili var; ön yazi ile karismis olabilir."))
    words = len(text.split())
    if words > 1200:
        f.append(("ONEMLI", f"~{words} kelime — cogu pazar icin uzun. 2 sayfayi asiyor olabilir."))
    elif words < 250:
        f.append(("ONEMLI", f"~{words} kelime — icerik ince: ya CV cok kisa ya da metnin bir kismi cikarilamadi."))
    return f, {"kelime": words}


def check_jd_overlap(text, jd_text):
    low = text.lower()
    tokens = re.findall(r"[a-zçğıöşüA-ZÇĞİÖŞÜ][\w\-\.\+#]{2,}", jd_text.lower())
    tokens = [t.strip(".,;:-") for t in tokens]
    freq = Counter(t for t in tokens if len(t) > 2 and t not in STOPWORDS and not t.isdigit())
    key_terms = [t for t, c in freq.most_common(60) if c >= 2][:25]
    missing = [t for t in key_terms if t not in low]
    covered = len(key_terms) - len(missing)
    f = []
    if key_terms:
        oran = covered / len(key_terms)
        sev = "KRITIK" if oran < 0.4 else ("ONEMLI" if oran < 0.65 else "INCE")
        f.append((sev, f"Ilandaki tekrar eden terimlerin %{oran*100:.0f}'i CV'de geciyor "
                       f"({covered}/{len(key_terms)})."))
        if missing:
            f.append(("BILGI", "CV'de gecmeyenler: " + ", ".join(missing[:15]) +
                              "  — YALNIZCA gercekten sahip olduklarini, kendi cumlelerinin "
                              "icinde ekle. Kopyalama ve yigma yapma."))
    return f, {"terim": len(key_terms), "karsilanan": covered, "eksik": missing[:15]}


# ---------- rapor ----------

ORDER = {"KRITIK": 0, "ONEMLI": 1, "INCE": 2, "BILGI": 3}
LABEL = {"KRITIK": "KRİTİK", "ONEMLI": "ÖNEMLİ", "INCE": "İNCE AYAR", "BILGI": "BİLGİ"}


def main():
    ap = argparse.ArgumentParser(description="CV dosyasinin ayristirma guvenligini test eder")
    ap.add_argument("dosya")
    ap.add_argument("--job-description", "-j", help="ilan metni dosyasi (txt)")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    path = Path(args.dosya)
    if not path.exists():
        sys.exit(f"Dosya yok: {path}")

    pages, meta = extract(path)
    text = "\n".join(pages)

    findings, total_chars = check_text_presence(pages, meta)
    stats = {"karakter": total_chars, "sayfa": len(pages)}

    if total_chars >= 200:
        for fn in (check_contact, check_sections, check_dates):
            f, s = fn(text)
            findings += f
            stats[fn.__name__.replace("check_", "")] = s
        findings += check_layout(pages, meta, path.suffix.lower())
        f, s = check_evidence(text); findings += f; stats["kanit"] = s
        f, s = check_language(text); findings += f; stats["dil"] = s
        if args.job_description:
            jd = Path(args.job_description).read_text(encoding="utf-8", errors="ignore")
            f, s = check_jd_overlap(text, jd); findings += f; stats["ilan"] = s

    findings.sort(key=lambda x: ORDER.get(x[0], 9))

    if args.json:
        print(json.dumps({"dosya": str(path), "istatistik": stats,
                          "bulgular": [{"seviye": s, "mesaj": m} for s, m in findings]},
                         ensure_ascii=False, indent=2))
        return

    print(f"\n=== PARSE KONTROLÜ: {path.name} ===")
    print(f"Çıkarılan metin: {total_chars} karakter, {len(pages)} sayfa")
    if "kanit" in stats and stats["kanit"]:
        k = stats["kanit"]
        print(f"Madde: {k.get('madde')} · rakam içeren: {k.get('rakamli')} (%{k.get('oran',0)*100:.0f})")
    if "dil" in stats:
        print(f"Kelime: {stats['dil']['kelime']}")
    if not findings:
        print("\nSorun bulunamadı — dosya ayrıştırma açısından temiz.")
        return
    cur = None
    for sev, msg in findings:
        if sev != cur:
            print(f"\n[{LABEL.get(sev, sev)}]")
            cur = sev
        print(f"  - {msg}")
    print("\nNot: Bu test ayrıştırılabilirliği ölçer, içerik kalitesini değil. "
          "İçerik için denetim rubriğini uygula.")


if __name__ == "__main__":
    main()
